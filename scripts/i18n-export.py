#!/usr/bin/env python3
"""Export the site's text to a spreadsheet for a translator.

    pip install openpyxl
    python3 scripts/i18n-export.py urh            # -> reallow-translations-urh.xlsx
    python3 scripts/i18n-export.py yo out.xlsx    # any locale; existing drafts are pre-filled

The translator fills in the language column and sends the file back; load it with
scripts/i18n-import.py. Keys are never edited — they tie each row to the code.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MESSAGES = Path(__file__).resolve().parent.parent / "src" / "lib" / "i18n" / "messages"

LANGUAGES = {"yo": "Yorùbá", "ha": "Hausa", "ig": "Igbo", "pcm": "Nigerian Pidgin", "urh": "Urhobo"}

# Where each group of text appears on the site, so the translator has some context.
WHERE = {
    "about": "About page",
    "agreement": "Tenancy / sale agreement page",
    "app": "Application page (someone applying for a property)",
    "apply": "“Apply for this property” button and form",
    "auth": "Sign in / sign up / verify email pages",
    "badge": "“Verified” badge next to names",
    "camera": "Taking a profile photo with the camera",
    "common": "Buttons and words used everywhere (Save, Cancel, Back…)",
    "contact": "Contact page",
    "dash": "Dashboard (the account home page)",
    "decision": "Landlord accepting / declining an applicant",
    "faq": "Help centre questions (.q) and answers (.a)",
    "faqTopic": "Help centre topic names",
    "feedback": "Rating a meeting or inspection afterwards",
    "feedbackTag": "Quick-pick comments when rating a meeting",
    "footer": "Bottom of every page",
    "form": "Form for listing a property",
    "furnishing": "Furnished / unfurnished options",
    "heard": "Sign-up question: “How did you hear about us?”",
    "help": "Help centre page",
    "home": "Home page",
    "hours": "Office opening hours",
    "id": "Verifying your identity (NIN / driver's licence)",
    "intent": "Sign-up question: “What are you using Reallow for?”",
    "listing": "A single property's page",
    "listingStatus": "Property status labels (Live, Rented, Sold…)",
    "listings": "Browse properties page and its filters",
    "meetings": "Meetings & inspections calendar",
    "menu": "Account menu (profile picture, top right)",
    "nav": "Top navigation bar",
    "newListing": "“List your property” page",
    "notif": "Notifications (bell icon) and their messages",
    "pay": "Choosing how to pay (wallet or card)",
    "people": "Home page photo captions (everyday Nigerians)",
    "profile": "Your profile / account details",
    "profilePage": "Someone's public profile page",
    "ptype": "Property types (flat, duplex, self-contain…)",
    "pwrule": "Password rules shown while typing a password",
    "receipt": "Payment receipt page",
    "report": "Reporting a property or person",
    "resubmit": "Re-submitting a listing that was rejected",
    "role": "Account type labels",
    "save": "“Save for later” button on a property",
    "settings": "Settings page",
    "terms": "Accepting the terms at sign-up",
    "ticketStatus": "Support ticket status labels",
    "tickets": "Support messages to Reallow",
    "tx": "Transactions / payments list",
    "txType": "Payment type labels",
    "upload": "Uploading photos and videos",
    "visit": "Property verification visit by Reallow staff",
    "wallet": "Reallow wallet and referral earnings",
}

INSTRUCTIONS = [
    "How to translate this file",
    "",
    "1. Fill in the {lang} column only. Leave the Key and English columns exactly as they are.",
    "2. Words in curly brackets, like {{count}}, {{name}} or {{year}}, are filled in by the website "
    "(a number, a name, a date…). Copy them into your translation unchanged — don't translate the "
    "word inside the brackets. You can move them to wherever they sit naturally in the sentence.",
    "3. Keep these names as they are: Reallow, NIN, BVN, WhatsApp, Facebook, Instagram, TikTok, "
    "Google, X (Twitter), city names (Warri, Lagos, Abuja…), and ₦ amounts.",
    "4. If you're not sure about a line, write your best attempt and add a note in the Notes column.",
    "5. Leave a cell empty if you want to skip it — the website will show the English for that line.",
    "6. Some lines are very similar (e.g. “1 photo” and “{{count}} photos”). That's on purpose: "
    "the first is used for one thing, the second for several.",
    "",
    "Rows already filled in are drafts — please check and correct them.",
]


def main() -> None:
    if len(sys.argv) < 2 or sys.argv[1] not in LANGUAGES:
        sys.exit(f"usage: i18n-export.py <{'|'.join(LANGUAGES)}> [out.xlsx]")
    locale = sys.argv[1]
    lang = LANGUAGES[locale]
    out = Path(sys.argv[2] if len(sys.argv) > 2 else f"reallow-translations-{locale}.xlsx")

    english = json.loads((MESSAGES / "en.json").read_text())
    existing = json.loads((MESSAGES / f"{locale}.json").read_text())

    wb = Workbook()
    guide = wb.active
    guide.title = "Instructions"
    guide.column_dimensions["A"].width = 110
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = guide.cell(row=i, column=1, value=line.format(lang=lang))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=14)

    sheet = wb.create_sheet("Translations")
    headers = ["Key", "Where it appears", "English", lang, "Notes"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="B5532D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    to_fill = PatternFill("solid", fgColor="FFF7E6")

    for key in sorted(english):
        namespace = key.split(".")[0]
        placeholders = sorted(set(re.findall(r"\{\w+\}", english[key])))
        note = f"Keep {', '.join(placeholders)}" if placeholders else ""
        sheet.append([key, WHERE.get(namespace, namespace), english[key], existing.get(key, ""), note])
        sheet.cell(row=sheet.max_row, column=4).fill = to_fill

    for col, width in zip("ABCDE", (34, 34, 60, 60, 30)):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    wb.active = 1

    wb.save(out)
    done = sum(1 for k in english if existing.get(k))
    print(f"{out}: {len(english)} lines, {done} already translated into {lang}")


if __name__ == "__main__":
    main()
