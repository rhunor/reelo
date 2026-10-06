#!/usr/bin/env python3
"""Load a translator's spreadsheet (made by scripts/i18n-export.py) back into the site.

    pip install openpyxl
    python3 scripts/i18n-import.py urh reallow-translations-urh.xlsx

Filled-in cells replace that language's text; empty cells are left alone (the site shows
English for them). Rows whose {placeholders} don't match the English are skipped and listed,
since a missing {count} or {name} would show up broken on the site.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

MESSAGES = Path(__file__).resolve().parent.parent / "src" / "lib" / "i18n" / "messages"
LOCALES = ("yo", "ha", "ig", "pcm", "urh")


def placeholders(text: str) -> set[str]:
    return set(re.findall(r"\{\w+\}", text))


def main() -> None:
    if len(sys.argv) != 3 or sys.argv[1] not in LOCALES:
        sys.exit(f"usage: i18n-import.py <{'|'.join(LOCALES)}> <file.xlsx>")
    locale, path = sys.argv[1], sys.argv[2]

    english = json.loads((MESSAGES / "en.json").read_text())
    target_path = MESSAGES / f"{locale}.json"
    target = json.loads(target_path.read_text())

    sheet = load_workbook(path, read_only=True)["Translations"]
    updated, skipped = 0, []
    for key, _where, _english, text, *_ in sheet.iter_rows(min_row=2, values_only=True):
        if not key or key not in english or text is None or not str(text).strip():
            continue
        text = str(text).strip()
        if placeholders(text) != placeholders(english[key]):
            skipped.append(key)
            continue
        if target.get(key) != text:
            target[key] = text
            updated += 1

    target_path.write_text(json.dumps(dict(sorted(target.items())), ensure_ascii=False, indent=2) + "\n")
    print(f"{locale}: {updated} lines updated, {len(target)}/{len(english)} translated")
    if skipped:
        print("Skipped — {placeholders} don't match the English:")
        for key in skipped:
            print(f"  {key}: needs {', '.join(sorted(placeholders(english[key]))) or 'none'}")


if __name__ == "__main__":
    main()
